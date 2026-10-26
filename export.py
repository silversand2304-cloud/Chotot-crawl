"""Dữ liệu theo từ khóa: bảng thống kê nhóm (dùng chung cho dashboard, lệnh search/export) và xuất
Excel/CSV. Không xuất mô tả tin, tên/số điện thoại người bán — chỉ những gì đã lưu trong DB."""
import csv
import io
import json
import os
import time
from datetime import datetime, timedelta, timezone

import analyzer
import config
import db
from normalizer import FLAG_LABEL, MARKET_VN, NO_ALERT, SKIP_LABEL, slug

DAY = 86400
VN_TZ = timezone(timedelta(hours=7))


def group_rows(conn, now, watch=None):
    """Thống kê giá từng nhóm trong cửa sổ WINDOW_DAYS ngày (lọc theo từ khóa nếu có)."""
    out = []
    for key, label, n in db.key_counts(conn, now - config.WINDOW_DAYS * DAY, watch=watch):
        rows = db.history(conn, key, now - config.WINDOW_DAYS * DAY)
        s = analyzer.stats([r[0] for r in rows], analyzer.adjusted_history(rows, key))
        m7 = analyzer.median_clean(db.prices(conn, key, now - config.TREND_DAYS * DAY))
        parts = key.split("|")
        flag = next((p for p in parts[2:] if p in NO_ALERT), None)
        out.append({
            "key": key, "label": label, "market": parts[1] if len(parts) > 1 else "",
            "market_vn": MARKET_VN.get(parts[1], parts[1]) if len(parts) > 1 else "", "n": n, "n_clean": s["n"],
            "median": s["median"], "p25": s["p25"], "p75": s["p75"], "m7": m7,
            "min": min((r[0] for r in rows), default=None), "max": max((r[0] for r in rows), default=None),
            "spread": s["spread"], "cut": s["cut"], "flag": FLAG_LABEL.get(flag),
            "dispersed": bool(s.get("dispersed")),
            "trend": (m7 - s["median"]) / s["median"] if m7 and s["median"] else None,
            "ready": s["n"] >= config.MIN_SAMPLES and not flag,  # đủ mẫu sau khi lọc ngoại lai
        })
    return out


def _dt(ts):
    return datetime.fromtimestamp(float(ts), VN_TZ).replace(tzinfo=None) if ts else None


def rows(conn, watch=None):
    """Mọi tin của 1 từ khóa (hoặc tất cả), mới nhất trước. Tin khớp 2 từ khóa xuất hiện ở cả hai."""
    sql = ("SELECT k.watch, k.key AS gkey, k.label AS glabel, l.*, a.kind AS alert, a.discount AS alert_discount "
           "FROM listing_keys k JOIN listings l USING (list_id) LEFT JOIN alerts a USING (list_id)")
    args = []
    if watch:
        sql += " WHERE k.watch = ?"
        args.append(watch)
    out = []
    for r in conn.execute(sql + " ORDER BY k.watch, l.first_seen DESC", args).fetchall():
        attrs = json.loads(r["attrs"]) if r["attrs"] else {}
        parts = (r["gkey"] or "").split("|")
        flag = next((p for p in parts[2:] if p in NO_ALERT), None)
        out.append({
            "Từ khóa": r["watch"] or "(chưa gắn từ khóa)",
            "Nhóm": r["glabel"],
            "Tình trạng / nguồn": MARKET_VN.get(parts[1], parts[1]) if len(parts) > 1 else "",
            "Dấu hiệu": FLAG_LABEL.get(flag, "đã sửa" if "repaired" in parts else ""),
            "Tiêu đề": r["title"],
            "Giá (đ)": r["price"],
            "Khu vực": r["region"],
            "Người bán": r["seller"],
            "Danh mục": attrs.get("category_name") or "",
            "Năm": attrs.get("year"),
            "Số km": attrs.get("mileage"),
            "Pin %": r["battery"],
            "Ngoại hình %": r["cosmetic"],
            "Đăng/đẩy lúc": _dt(r["posted_at"]),
            "Đăng gốc lúc": _dt(r["orig_posted_at"]),
            "Thấy lần đầu": _dt(r["first_seen"]),
            "Còn đăng lúc": _dt(r["last_seen"]),
            "Đã gỡ lúc": _dt(r["gone_at"]),
            "Cảnh báo": {"deal": "deal", "scam_risk": "nghi lừa đảo"}.get(r["alert"], ""),
            "Rẻ hơn trung vị": r["alert_discount"],
            "Link": r["url"],
            "Mã tin": r["list_id"],
            "Mã nhóm": r["gkey"],
        })
    return out


def skipped_records(conn, watch=None):
    """Tin bị loại kèm lý do (không có mô tả/tên/số điện thoại người bán)."""
    out = []
    for r in db.skipped_rows(conn, watch=watch, limit=100_000):
        out.append({
            "Từ khóa": r["watch"] or "(chưa gắn từ khóa)",
            "Lý do": SKIP_LABEL.get(r["reason"], r["reason"]),
            "Chi tiết": r["detail"] or "",
            "Tiêu đề": r["title"],
            "Giá (đ)": r["price"],
            "Khu vực": r["region"],
            "Người bán": r["seller"],
            "Danh mục": r["category"] or "",
            "Thấy lần đầu": _dt(r["first_seen"]),
            "Thấy lần cuối": _dt(r["last_seen"]),
            "Link": r["url"],
            "Mã tin": r["list_id"],
        })
    return out


def summary(conn, now, watch=None):
    out = []
    names = [watch] if watch else [w["name"] for w in config.WATCHLIST]
    for name in names:
        for g in group_rows(conn, now, watch=name):
            out.append({
                "Từ khóa": name, "Nhóm": g["label"], "Tình trạng / nguồn": g["market_vn"],
                f"Số tin {config.WINDOW_DAYS} ngày": g["n"], "Trung vị (đ)": g["median"],
                f"Trung vị {config.TREND_DAYS} ngày (đ)": g["m7"], "Xu hướng": g["trend"],
                "P25 (đ)": g["p25"], "P75 (đ)": g["p75"], "Thấp nhất (đ)": g["min"], "Cao nhất (đ)": g["max"],
                "Độ phân tán": g["spread"],
                "Báo deal khi rẻ hơn": None if (g["flag"] or g["dispersed"] or not g["ready"]) else g["cut"],
                "Ghi chú": (f"không báo ({g['flag']})" if g["flag"] else "giá quá phân tán" if g["dispersed"]
                            else "" if g["ready"] else "chưa đủ mẫu"),
                "Mã nhóm": g["key"],
            })
    return out


_MONEY = {"Giá (đ)", "Trung vị (đ)", "P25 (đ)", "P75 (đ)", "Thấp nhất (đ)", "Cao nhất (đ)"}
_PCT = {"Xu hướng", "Độ phân tán", "Báo deal khi rẻ hơn", "Rẻ hơn trung vị"}
_WIDTH = {"Tiêu đề": 60, "Nhóm": 34, "Từ khóa": 22, "Link": 14, "Mã nhóm": 30, "Ghi chú": 20, "Lý do": 34, "Chi tiết": 36}


def _sheet(ws, records, empty_note):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not records:
        ws.append([empty_note])
        return
    cols = list(records[0])
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF8")
    for rec in records:
        ws.append([rec[c] for c in cols])
    for i, col in enumerate(cols, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = _WIDTH.get(col, max(11, min(len(col) + 3, 20)))
        fmt = ("#,##0" if col in _MONEY or col.startswith("Trung vị") else "0.0%" if col in _PCT
               else "dd/mm/yyyy hh:mm" if col.endswith("lúc") or col.startswith("Thấy") else None)
        for cell in ws[letter][1:]:
            if fmt:
                cell.number_format = fmt
            if col == "Link" and cell.value:
                cell.hyperlink, cell.value, cell.style = cell.value, "Mở tin", "Hyperlink"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def to_xlsx(conn, now, watch=None):
    """File Excel (bytes): sheet "Tổng hợp" (giá theo nhóm, 30 ngày) + "Tin đăng" (mọi tin đã thu) +
    "Tin bỏ qua" (tin bị loại kèm lý do) + "Ghi chú"."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    _sheet(ws, summary(conn, now, watch), "Chưa có dữ liệu.")
    _sheet(wb.create_sheet("Tin đăng"), rows(conn, watch), "Chưa có tin nào.")
    _sheet(wb.create_sheet("Tin bỏ qua"), skipped_records(conn, watch),
           "Chưa có tin bỏ qua nào (danh sách được ghi từ lúc collector chạy bản có tab Tin bỏ qua).")
    info = wb.create_sheet("Ghi chú")
    for line in (f"Xuất lúc {datetime.now(VN_TZ):%d/%m/%Y %H:%M} từ {os.path.basename(config.DB_PATH)}",
                 f"Từ khóa: {watch or 'tất cả'}",
                 f"Tổng hợp: giá chào bán {config.WINDOW_DAYS} ngày gần nhất, trung vị sau khi lọc ngoại lai (IQR).",
                 "Giá là giá người bán rao, không phải giá chốt. Tin biến mất ≠ chắc chắn đã bán.",
                 f"Tin bỏ qua: tin bị loại khỏi bảng giá (thu mua, ngoài khoảng giá, thiếu chữ từ khóa…), "
                 f"giữ {config.SKIP_KEEP_DAYS} ngày kể từ lần cuối thấy.",
                 "Không có mô tả tin, tên hay số điện thoại người bán (không thu thập)."):
        info.append([line])
    info.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv(conn, watch=None, skipped=False):
    """CSV UTF-8 có BOM (Excel đọc đúng tiếng Việt), dấu phân cách ';' cho Excel cài đặt Việt Nam.
    skipped=True: xuất danh sách tin bỏ qua thay cho tin đăng."""
    records = skipped_records(conn, watch) if skipped else rows(conn, watch)
    buf = io.StringIO()
    if records:
        w = csv.DictWriter(buf, fieldnames=list(records[0]), delimiter=";")
        w.writeheader()
        for r in records:
            w.writerow({k: (v.strftime("%d/%m/%Y %H:%M") if isinstance(v, datetime) else v) for k, v in r.items()})
    return ("﻿" + buf.getvalue()).encode("utf-8")


def filename(watch, ext, suffix=None):
    base = slug(watch) if watch else "tat-ca-tu-khoa"
    return f"{base}{'_' + suffix if suffix else ''}_{time.strftime('%Y%m%d_%H%M')}.{ext}"


def save_xlsx(conn, now, watch=None):
    os.makedirs(config.EXPORT_DIR, exist_ok=True)
    path = os.path.join(config.EXPORT_DIR, filename(watch, "xlsx"))
    with open(path, "wb") as f:
        f.write(to_xlsx(conn, now, watch))
    return path
